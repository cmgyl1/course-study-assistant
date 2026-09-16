"""核心引擎冒烟测试：验证文档转换管线（txt/md/docx/pdf）可用。"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from engine.config import ensure_data_dirs
from engine.converter import convert_file, ConversionError

def main() -> int:
    ensure_data_dirs()
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        # 1) txt
        (tmp / "sample.txt").write_text("第一章 计算机网络概述\n\n这是文本内容。", encoding="utf-8")
        # 2) md
        (tmp / "sample.md").write_text("# 标题\n\n正文段落。", encoding="utf-8")
        # 3) docx
        from docx import Document
        doc = Document()
        doc.add_heading("计算机网络概述", level=1)
        doc.add_paragraph("这是用 python-docx 生成的测试文档。")
        doc.add_paragraph("第二章 物理层")
        doc.save(str(tmp / "sample.docx"))
        # 4) pdf
        import fitz
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "第一章 计算机网络概述")
        page.insert_text((72, 100), "这是 PDF 测试内容。")
        pdf.save(str(tmp / "sample.pdf"))
        pdf.close()

        ok = True
        for f in ["sample.txt", "sample.md", "sample.docx", "sample.pdf"]:
            try:
                out = convert_file(tmp / f, tmp / "out")
                text = out.read_text(encoding="utf-8")
                print(f"[OK] {f} -> {out.name} ({len(text)} 字符)")
                print("     preview:", text.strip().replace("\n", " | ")[:80])
            except ConversionError as e:
                print(f"[FAIL] {f}: {e}")
                ok = False
        return 0 if ok else 1

if __name__ == "__main__":
    sys.exit(main())
